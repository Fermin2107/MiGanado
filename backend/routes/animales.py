import io
import re
from datetime import date as date_type
from flask import Blueprint, request, jsonify, send_file
from flask_jwt_extended import jwt_required, get_jwt_identity
from sqlalchemy import or_, text, func
from sqlalchemy.exc import IntegrityError
from models import db, Animal, Raza, EventoAnimal

animales_bp = Blueprint("animales", __name__)

ANIMAL_FIELDS = [
    "rp", "caravana_rfid", "hba", "nombre", "sexo", "fecha_nac", "nacimiento", "color",
    "padre", "madre", "padre_id", "madre_id", "abuelo_paterno", "abuelo_materno",
    "familia", "f", "tamano", "pezunas", "articulacion", "ap_delanteros", "ap_traseros",
    "curv_garrones", "apert_posterior", "ubres_pezones", "forma_testicular", "desplazamiento",
    "clase", "impresion_general", "musculatura", "anchura", "costilla", "docilidad",
    "valoracion", "observaciones", "premios",
    "epd_nac", "epd_dest", "epd_leche", "epd_18m", "epd_pa_v", "epd_ce",
    "epd_aob", "epd_egs", "epd_marb",
    "val_14m", "val_18m", "val_ternero", "val_adulto",
]

DATE_FIELDS  = {"fecha_nac"}
FLOAT_FIELDS = {
    "pezunas", "articulacion", "clase", "valoracion",
    "epd_nac", "epd_dest", "epd_leche", "epd_18m", "epd_pa_v",
    "epd_ce", "epd_aob", "epd_egs", "epd_marb",
}
INT_FIELDS   = {"padre_id", "madre_id"}

SORTABLE_COLS = {
    "rp", "nombre", "sexo", "fecha_nac", "padre", "madre",
    "valoracion", "clase", "nacimiento", "color",
}


def _get_user_raza(raza_id):
    user_id = int(get_jwt_identity())
    return Raza.query.filter_by(id=raza_id, user_id=user_id).first_or_404()


def _apply_data(animal, data):
    for field in ANIMAL_FIELDS:
        if field not in data:
            continue
        val = data[field]
        if val == "" or val is None:
            val = None
        if val is not None and field in DATE_FIELDS:
            from datetime import date
            val = date.fromisoformat(val) if isinstance(val, str) else val
        if val is not None and field in FLOAT_FIELDS:
            try:
                val = float(val)
            except (TypeError, ValueError):
                val = None
        if val is not None and field in INT_FIELDS:
            try:
                val = int(val)
            except (TypeError, ValueError):
                val = None
        setattr(animal, field, val)


def _build_rp_order(dir_suffix):
    return [
        text("CASE WHEN rp ~ '^[0-9]+$' THEN 0 ELSE 1 END ASC"),
        text(f"(CASE WHEN rp ~ '^[0-9]+$' THEN rp::int END) {dir_suffix} NULLS LAST"),
        text(f"rp {dir_suffix} NULLS LAST"),
    ]


# ── Árbol genealógico (helpers) ───────────────────────────────────────────────

def _text_node(nombre, rp=None, sexo=None):
    """Nodo externo: texto libre, no tiene id en el sistema."""
    if not nombre:
        return None
    return {
        "id": None, "nombre": nombre, "rp": rp,
        "sexo": sexo, "externo": True,
        "padre": None, "madre": None,
    }


def _real_node(animal, depth=1):
    """Nodo real (animal en el sistema). depth=1 → incluye sus padres inmediatos."""
    if animal is None:
        return None
    node = {
        "id": animal.id, "nombre": animal.nombre,
        "rp": animal.rp,   "sexo": animal.sexo,
        "externo": False,
    }
    if depth <= 0:
        node["padre"] = None
        node["madre"] = None
        return node

    # Padre
    if animal.padre_rel:
        node["padre"] = _real_node(animal.padre_rel, depth - 1)
    elif animal.padre:
        stub = _text_node(animal.padre)
        stub["padre"] = _text_node(animal.abuelo_paterno)
        node["padre"] = stub
    else:
        node["padre"] = None

    # Madre
    if animal.madre_rel:
        node["madre"] = _real_node(animal.madre_rel, depth - 1)
    elif animal.madre:
        stub = _text_node(animal.madre)
        stub["padre"] = _text_node(animal.abuelo_materno)
        node["madre"] = stub
    else:
        node["madre"] = None

    return node


# ── Listado y creación por raza ───────────────────────────────────────────────

@animales_bp.get("/api/razas/<int:raza_id>/animales")
@jwt_required()
def listar_animales(raza_id):
    _get_user_raza(raza_id)
    page     = request.args.get("page", 1, type=int)
    per_page = min(request.args.get("per_page", 25, type=int), 100)
    sexo     = request.args.get("sexo", "").strip()
    search   = request.args.get("search", "").strip()
    sort_by  = request.args.get("sort_by", "rp").strip()
    sort_dir = request.args.get("sort_dir", "asc").strip().lower()

    if sort_by not in SORTABLE_COLS:
        sort_by = "rp"
    dir_suffix = "DESC" if sort_dir == "desc" else "ASC"

    q = Animal.query.filter_by(raza_id=raza_id)
    if sexo:
        q = q.filter(Animal.sexo == sexo)
    if search:
        like = f"%{search}%"
        q = q.filter(or_(
            Animal.rp.ilike(like), Animal.nombre.ilike(like),
            Animal.padre.ilike(like), Animal.madre.ilike(like),
        ))

    if sort_by == "rp":
        q = q.order_by(*_build_rp_order(dir_suffix))
    elif sort_by == "fecha_nac":
        col = Animal.fecha_nac
        q = q.order_by(col.desc() if sort_dir == "desc" else col.asc())
    elif sort_by == "valoracion":
        col = Animal.valoracion
        q = q.order_by(col.desc().nullslast() if sort_dir == "desc" else col.asc().nullslast())
    else:
        col = getattr(Animal, sort_by, None)
        if col is not None:
            q = q.order_by(col.desc() if sort_dir == "desc" else col.asc())
        else:
            q = q.order_by(*_build_rp_order("ASC"))

    pag = q.paginate(page=page, per_page=per_page, error_out=False)
    return jsonify({
        "animales":    [a.to_dict() for a in pag.items],
        "total":       pag.total,
        "page":        page,
        "per_page":    per_page,
        "total_pages": pag.pages,
    }), 200


@animales_bp.post("/api/razas/<int:raza_id>/animales")
@jwt_required()
def crear_animal(raza_id):
    _get_user_raza(raza_id)
    data = request.get_json() or {}

    if not (data.get("rp") or "").strip():
        return jsonify({"error": "El RP es obligatorio"}), 400
    if not (data.get("sexo") or "").strip():
        return jsonify({"error": "El sexo es obligatorio"}), 400

    animal = Animal(raza_id=raza_id)
    _apply_data(animal, data)
    db.session.add(animal)
    try:
        db.session.commit()
    except IntegrityError:
        db.session.rollback()
        rp = (data.get("rp") or "").strip()
        return jsonify({"error": f"Ya existe un animal con el RP '{rp}' en esta raza"}), 409
    return jsonify(animal.to_dict()), 201


# ── Ficha, edición y eliminación ─────────────────────────────────────────────

@animales_bp.get("/api/animales/<int:animal_id>")
@jwt_required()
def get_animal(animal_id):
    user_id = int(get_jwt_identity())
    animal = Animal.query.join(Raza).filter(
        Animal.id == animal_id, Raza.user_id == user_id
    ).first_or_404()

    # Hijos por ID directo
    hijos_por_id = Animal.query.join(Raza).filter(
        Raza.user_id == user_id,
        Animal.id != animal.id,
        or_(Animal.padre_id == animal.id, Animal.madre_id == animal.id),
    ).all()

    # Hijos por comparación de texto (fallback)
    hijos_id_set = {h.id for h in hijos_por_id}
    text_conditions = []
    if animal.rp:
        text_conditions += [Animal.padre == animal.rp, Animal.madre == animal.rp]
    if animal.nombre:
        nl = animal.nombre.lower()
        text_conditions += [
            func.lower(Animal.padre) == nl,
            func.lower(Animal.madre) == nl,
        ]
    hijos_por_texto = []
    if text_conditions:
        hijos_por_texto = Animal.query.join(Raza).filter(
            Raza.user_id == user_id,
            Animal.id != animal.id,
            or_(*text_conditions),
        ).all()

    all_hijos = hijos_por_id + [h for h in hijos_por_texto if h.id not in hijos_id_set]
    all_hijos = all_hijos[:50]

    data = animal.to_dict(include_eventos=True)
    data["hijos"] = [h.to_dict() for h in all_hijos]
    return jsonify(data), 200


@animales_bp.put("/api/animales/<int:animal_id>")
@jwt_required()
def editar_animal(animal_id):
    user_id = int(get_jwt_identity())
    animal = Animal.query.join(Raza).filter(
        Animal.id == animal_id, Raza.user_id == user_id
    ).first_or_404()
    data = request.get_json() or {}
    _apply_data(animal, data)
    try:
        db.session.commit()
    except IntegrityError:
        db.session.rollback()
        rp = (data.get("rp") or "").strip()
        return jsonify({"error": f"Ya existe un animal con el RP '{rp}' en esta raza"}), 409
    return jsonify(animal.to_dict()), 200


@animales_bp.delete("/api/animales/<int:animal_id>")
@jwt_required()
def eliminar_animal(animal_id):
    user_id = int(get_jwt_identity())
    animal = Animal.query.join(Raza).filter(
        Animal.id == animal_id, Raza.user_id == user_id
    ).first_or_404()
    db.session.delete(animal)
    db.session.commit()
    return jsonify({"message": "Animal eliminado"}), 200


# ── Árbol genealógico ─────────────────────────────────────────────────────────

@animales_bp.get("/api/animales/<int:animal_id>/arbol")
@jwt_required()
def arbol_genealogico(animal_id):
    user_id = int(get_jwt_identity())
    animal = Animal.query.join(Raza).filter(
        Animal.id == animal_id, Raza.user_id == user_id
    ).first_or_404()

    root = {
        "id": animal.id, "nombre": animal.nombre,
        "rp": animal.rp,   "sexo": animal.sexo,
        "externo": False,
    }

    # Padre (con sus padres = abuelos paternos)
    if animal.padre_rel:
        padre_node = _real_node(animal.padre_rel, depth=1)
    elif animal.padre:
        padre_node = _text_node(animal.padre)
        padre_node["padre"] = _text_node(animal.abuelo_paterno)
        padre_node["madre"] = None
    else:
        padre_node = None

    # Madre (con sus padres = abuelos maternos)
    if animal.madre_rel:
        madre_node = _real_node(animal.madre_rel, depth=1)
    elif animal.madre:
        madre_node = _text_node(animal.madre)
        madre_node["padre"] = _text_node(animal.abuelo_materno)
        madre_node["madre"] = None
    else:
        madre_node = None

    return jsonify({
        "animal": root,
        "padre":  padre_node,
        "madre":  madre_node,
    }), 200


# ── Exportar Excel ────────────────────────────────────────────────────────────

# Orden y headers de columnas para el Excel
_EXPORT_COLS = [
    ("rp",               "RP"),
    ("caravana_rfid",    "Caravana RFID"),
    ("hba",              "HBA"),
    ("nombre",           "Nombre"),
    ("sexo",             "Sexo"),
    ("fecha_nac",        "Fecha Nacimiento"),
    ("color",            "Color"),
    ("padre",            "Padre"),
    ("madre",            "Madre"),
    ("abuelo_paterno",   "Abuelo Paterno"),
    ("abuelo_materno",   "Abuelo Materno"),
    ("familia",          "Familia"),
    ("tamano",           "Tamaño"),
    ("pezunas",          "Pezuñas"),
    ("articulacion",     "Articulación"),
    ("clase",            "Clase"),
    ("valoracion",       "Valoración"),
    ("epd_nac",          "DEP Peso Nacer"),
    ("epd_dest",         "DEP Peso Destete"),
    ("epd_leche",        "DEP Hab. Materna/Leche"),
    ("epd_18m",          "DEP Peso 18M"),
    ("epd_pa_v",         "DEP Peso Adulto Vaca"),
    ("epd_ce",           "DEP Circ. Escrotal"),
    ("epd_aob",          "DEP AOB"),
    ("epd_egs",          "DEP EGS"),
    ("epd_marb",         "DEP Marbling"),
    ("val_14m",          "Valoración 14M"),
    ("val_18m",          "Valoración 18M"),
    ("val_ternero",      "Valoración Ternero"),
    ("val_adulto",       "Valoración Adulto"),
    ("observaciones",    "Observaciones"),
    ("premios",          "Premios"),
]

_NUMERIC_FIELDS = {
    "pezunas", "articulacion", "clase", "valoracion",
    "epd_nac", "epd_dest", "epd_leche", "epd_18m", "epd_pa_v",
    "epd_ce", "epd_aob", "epd_egs", "epd_marb",
}

_HEADER_GREEN = "1B4332"   # tema primario sin #


@animales_bp.post("/api/razas/<int:raza_id>/animales/exportar")
@jwt_required()
def exportar_animales(raza_id):
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter

    raza = _get_user_raza(raza_id)
    body     = request.get_json() or {}
    sexo     = body.get("sexo",     "").strip()
    search   = body.get("search",   "").strip()
    sort_by  = body.get("sort_by",  "rp").strip()
    sort_dir = body.get("sort_dir", "asc").strip().lower()

    if sort_by not in SORTABLE_COLS:
        sort_by = "rp"
    dir_suffix = "DESC" if sort_dir == "desc" else "ASC"

    q = Animal.query.filter_by(raza_id=raza_id)
    if sexo:
        q = q.filter(Animal.sexo == sexo)
    if search:
        like = f"%{search}%"
        q = q.filter(or_(
            Animal.rp.ilike(like), Animal.nombre.ilike(like),
            Animal.padre.ilike(like), Animal.madre.ilike(like),
        ))

    if sort_by == "rp":
        q = q.order_by(*_build_rp_order(dir_suffix))
    elif sort_by == "fecha_nac":
        col = Animal.fecha_nac
        q = q.order_by(col.desc() if sort_dir == "desc" else col.asc())
    elif sort_by == "valoracion":
        col = Animal.valoracion
        q = q.order_by(col.desc().nullslast() if sort_dir == "desc" else col.asc().nullslast())
    else:
        col = getattr(Animal, sort_by, None)
        if col is not None:
            q = q.order_by(col.desc() if sort_dir == "desc" else col.asc())
        else:
            q = q.order_by(*_build_rp_order("ASC"))

    animales = q.all()

    # ── Construir workbook ────────────────────────────────────────────────────
    wb = openpyxl.Workbook()
    ws = wb.active
    # Nombre de hoja: sanitizar caracteres no permitidos en Excel
    safe_name = re.sub(r'[\\/*?:\[\]]', '_', raza.nombre)[:31]
    ws.title = safe_name or "Animales"

    header_fill = PatternFill("solid", fgColor=_HEADER_GREEN)
    header_font = Font(color="FFFFFF", bold=True, name="Calibri", size=11)
    header_align = Alignment(horizontal="center", vertical="center", wrap_text=False)

    headers = [h for _, h in _EXPORT_COLS]
    ws.append(headers)

    for col_idx, cell in enumerate(ws[1], start=1):
        cell.fill  = header_fill
        cell.font  = header_font
        cell.alignment = header_align

    ws.freeze_panes = "A2"

    # ── Filas de datos ────────────────────────────────────────────────────────
    date_fmt  = "DD/MM/YYYY"
    num_fmt   = "0.##"      # hasta 2 decimales, sin ceros innecesarios

    for animal in animales:
        row = []
        for field, _ in _EXPORT_COLS:
            val = getattr(animal, field, None)
            row.append(val)
        ws.append(row)

        data_row = ws.max_row
        for col_idx, (field, _) in enumerate(_EXPORT_COLS, start=1):
            cell = ws.cell(row=data_row, column=col_idx)
            if field == "fecha_nac":
                if isinstance(cell.value, date_type):
                    cell.number_format = date_fmt
                else:
                    cell.value = None
            elif field in _NUMERIC_FIELDS:
                if cell.value is not None:
                    try:
                        cell.value = float(cell.value)
                        cell.number_format = num_fmt
                    except (TypeError, ValueError):
                        cell.value = None

    # ── Autoajuste de ancho ───────────────────────────────────────────────────
    for col_idx, (field, header) in enumerate(_EXPORT_COLS, start=1):
        col_letter = get_column_letter(col_idx)
        max_len = len(header)
        for row in ws.iter_rows(min_row=2, min_col=col_idx, max_col=col_idx):
            for cell in row:
                if cell.value is not None:
                    display = str(cell.value)
                    if isinstance(cell.value, date_type):
                        display = "DD/MM/AAAA"
                    max_len = max(max_len, len(display))
        ws.column_dimensions[col_letter].width = min(max_len + 4, 60)

    # ── Respuesta ─────────────────────────────────────────────────────────────
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)

    today = date_type.today().strftime("%Y%m%d")
    safe_raza = re.sub(r'[^\w\-]', '_', raza.nombre)
    filename = f"{safe_raza}_animales_{today}.xlsx"

    return send_file(
        buf,
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        as_attachment=True,
        download_name=filename,
    )


# ── Importar EPDs desde Excel ─────────────────────────────────────────────────

# Mapa: header del Excel → campo del modelo
_EPD_COLUMN_MAP = {
    'dep peso nacer':           'epd_nac',
    'dep peso destete':         'epd_dest',
    'dep peso 18m':             'epd_18m',
    'dep peso adulto vaca':     'epd_pa_v',
    'dep circ. escrotal':       'epd_ce',
    'dep hab. materna / leche': 'epd_leche',
    'dep hab. materna/leche':   'epd_leche',
    'dep aob':                  'epd_aob',
    'dep egs / grasa':          'epd_egs',
    'dep egs/grasa':            'epd_egs',
    'dep marbling':             'epd_marb',
}

# Display name para reportar valores inválidos
_EPD_DISPLAY = {v: k.title() for k, v in _EPD_COLUMN_MAP.items()}


@animales_bp.post("/api/razas/<int:raza_id>/animales/importar_epds")
@jwt_required()
def importar_epds(raza_id):
    import pandas as pd
    from sqlalchemy.exc import OperationalError

    _get_user_raza(raza_id)

    if 'archivo' not in request.files or not request.files['archivo'].filename:
        return jsonify({"error": "Se requiere un archivo Excel"}), 400

    archivo = request.files['archivo']
    ext = archivo.filename.rsplit('.', 1)[-1].lower() if '.' in archivo.filename else ''
    if ext not in ('xlsx', 'xls'):
        return jsonify({"error": "Solo se aceptan archivos .xlsx o .xls"}), 400
    if ext == 'xls':
        return jsonify({"error": "El formato .xls no está soportado. Abrí el archivo en Excel y guardalo como .xlsx (Excel Workbook)."}), 400

    buf = io.BytesIO(archivo.read())

    # ── Detectar fila de encabezado ────────────────────────────────────────────
    try:
        df_raw = pd.read_excel(buf, header=None, engine='openpyxl')
    except Exception as e:
        return jsonify({"error": f"No se pudo leer el archivo: {str(e)}"}), 400

    header_row_idx = None
    for i, row in df_raw.iterrows():
        for val in row:
            if pd.notna(val) and str(val).strip().upper() == 'RP':
                header_row_idx = i
                break
        if header_row_idx is not None:
            break

    if header_row_idx is None:
        return jsonify({"error": "No se encontró la columna 'RP' en el archivo. Verificá que el Excel tenga una columna con ese encabezado exacto."}), 400

    # ── Releer con header correcto, todo como string para limpieza manual ──────
    buf.seek(0)
    try:
        df = pd.read_excel(buf, header=header_row_idx, engine='openpyxl', dtype=str)
    except Exception as e:
        return jsonify({"error": f"Error leyendo los datos: {str(e)}"}), 400

    # Normalizar nombres de columna
    df.columns = [str(c).strip() for c in df.columns]
    if 'RP' not in df.columns:
        # buscar case-insensitive
        rp_col = next((c for c in df.columns if c.upper() == 'RP'), None)
        if rp_col is None:
            return jsonify({"error": "Columna 'RP' no encontrada tras parsear el encabezado."}), 400
        df = df.rename(columns={rp_col: 'RP'})

    # Mapear columnas EPD presentes en el archivo
    col_to_field = {}
    for col in df.columns:
        key = col.strip().lower()
        if key in _EPD_COLUMN_MAP:
            col_to_field[col] = _EPD_COLUMN_MAP[key]

    if not col_to_field:
        return jsonify({"error": "No se encontraron columnas DEP reconocidas. Revisá que los encabezados coincidan con los nombres esperados."}), 400

    # ── Procesar filas ─────────────────────────────────────────────────────────
    actualizados      = 0
    no_encontrados    = []
    valores_invalidos = []
    total_filas       = 0
    batch_count       = 0
    BATCH_SIZE        = 50

    for _, row in df.iterrows():
        rp_raw = row.get('RP')
        if pd.isna(rp_raw) if isinstance(rp_raw, float) else (rp_raw is None or str(rp_raw).strip() == '' or str(rp_raw).strip().lower() == 'nan'):
            continue

        # Limpiar RP: "123.0" → "123"
        rp_str = str(rp_raw).strip()
        if rp_str.endswith('.0') and rp_str[:-2].isdigit():
            rp_str = rp_str[:-2]

        if not rp_str:
            continue

        total_filas += 1

        animal = Animal.query.filter_by(raza_id=raza_id, rp=rp_str).first()
        if not animal:
            no_encontrados.append(rp_str)
            continue

        updated_any = False
        for col, field in col_to_field.items():
            raw = row.get(col)
            if raw is None or (isinstance(raw, float) and pd.isna(raw)):
                continue
            val_str = str(raw).strip()
            if val_str.lower() in ('nan', '', 'none'):
                continue
            val_str = val_str.replace(',', '.')
            try:
                setattr(animal, field, float(val_str))
                updated_any = True
            except (ValueError, TypeError):
                display_col = next(
                    (k for k, v in _EPD_COLUMN_MAP.items() if v == field),
                    field
                ).title()
                valores_invalidos.append({
                    "rp": rp_str,
                    "columna": display_col,
                    "valor": str(raw),
                })

        if updated_any:
            actualizados += 1

        batch_count += 1
        if batch_count >= BATCH_SIZE:
            try:
                db.session.commit()
            except OperationalError:
                db.session.rollback()
                try:
                    db.session.commit()
                except Exception:
                    db.session.rollback()
            except Exception:
                db.session.rollback()
            batch_count = 0

    # Commit final
    if batch_count > 0:
        try:
            db.session.commit()
        except OperationalError:
            db.session.rollback()
            try:
                db.session.commit()
            except Exception:
                db.session.rollback()
        except Exception:
            db.session.rollback()

    return jsonify({
        "actualizados":           actualizados,
        "no_encontrados":         no_encontrados,
        "valores_invalidos":      valores_invalidos,
        "total_filas_procesadas": total_filas,
    }), 200


# ── Búsqueda por RFID ─────────────────────────────────────────────────────────

@animales_bp.get("/api/animales/rfid/<string:codigo>")
@jwt_required()
def buscar_por_rfid(codigo):
    user_id = int(get_jwt_identity())
    codigo_clean = codigo.strip()

    # Búsqueda exacta (string tal cual)
    animal = Animal.query.join(Raza).filter(
        Animal.caravana_rfid == codigo_clean, Raza.user_id == user_id
    ).first()

    # Fallback: si el código es todo dígitos, probar sin ceros a la izquierda
    # (el lector puede mandar "001234" pero el sistema tiene "1234")
    if animal is None and codigo_clean.isdigit():
        sin_ceros = str(int(codigo_clean))
        if sin_ceros != codigo_clean:
            animal = Animal.query.join(Raza).filter(
                Animal.caravana_rfid == sin_ceros, Raza.user_id == user_id
            ).first()

    if animal is None:
        return jsonify({
            "error": f"No se encontró ningún animal con la caravana '{codigo_clean}'"
        }), 404

    return jsonify(animal.to_dict()), 200
